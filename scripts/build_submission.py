"""Create reproducible Korean submission files from the checked-in analysis data.

The script builds a four-section A4 report with charts, equations, and measured
model errors. It is intended to run from the repository root.
"""
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd

OUT = Path("submission")
OUT.mkdir(exist_ok=True, parents=True)
FIG = OUT / "assets"
FIG.mkdir(exist_ok=True)
FONT = "NanumSquare"

font_path = Path("/usr/share/fonts/truetype/nanum/NanumSquareR.ttf")
if font_path.exists():
    font_manager.fontManager.addfont(str(font_path))
else:
    FONT = "DejaVu Sans"
plt.rcParams.update({"font.family": FONT, "axes.unicode_minus": False, "font.size": 10})

df = pd.read_csv("data/SeoulBikeData.csv", encoding="utf-8-sig")
df = df[df["Functioning Day"].eq("Yes")].copy()
metrics = pd.read_csv("results/metrics.csv")
coeff = pd.read_csv("results/coefficients.csv")
examples = pd.read_csv("results/example_test_day.csv")
assert len(df) == 8465 and len(examples) == 24
hours = np.arange(24)
means = df.groupby("Hour")["Rented Bike Count"].mean().reindex(hours).to_numpy()
testm = metrics.query("split == 'test'").set_index("model")
mnames = ["mean", "harmonic_24h", "harmonic_24h_12h", "harmonic_weekend"]
c = {k: coeff.query("model == @k")["coefficient"].to_numpy() for k in mnames}
def features(t, wk, key):
    p=2*np.pi*np.asarray(t)/24
    mat=np.column_stack([np.ones_like(p),np.sin(p),np.cos(p),
                         np.sin(2*p),np.cos(2*p)])
    if key=="harmonic_24h":return mat[:,:3]
    if key=="harmonic_24h_12h":return mat
    return np.column_stack([mat,mat*np.asarray(wk)[:,None]])

def formula(lines, filename):
    height=max(0.48,0.52*len(lines))
    fig=plt.figure(figsize=(9.6,height))
    for i,s in enumerate(lines):
        fig.text(.015,1-(i+.65)/len(lines),s,fontsize=14.2,color="#17324d")
    fig.savefig(FIG/filename,dpi=210,transparent=True,bbox_inches="tight",pad_inches=.05)
    plt.close(fig)

formula([r"$\hat{y}_1(t)=a_0+a_1\sin(2\pi t/24)+b_1\cos(2\pi t/24)$"],"eq1.png")
formula([r"$\hat{y}_2(t)=\hat{y}_1(t)+a_2\sin(4\pi t/24)+b_2\cos(4\pi t/24)$"],"eq2.png")
formula([r"$\hat{\beta}=\arg\min_{\beta}\Vert X\beta-y\Vert_2^2,\quad X^TX\beta=X^Ty$"],"least.png")
formula([r"$\mathrm{MAE}=\frac{1}{n}\sum_{i=1}^n|y_i-\hat{y}_i|,\quad \mathrm{RMSE}=\sqrt{\frac{1}{n}\sum_{i=1}^n(y_i-\hat{y}_i)^2}$"],"error.png")

NAVY="#17324d";BLUE="#2563a6";TEAL="#007f78";ORANGE="#bb631f"
def axis(ax,title):
    ax.set(title=title,xlabel="시각 (시)",ylabel="대여 건수/시간")
    ax.set_xticks(range(0,24,2));ax.set_xlim(-.5,23.5)
    ax.grid(axis="y",alpha=.22)
    ax.spines[["right","top"]].set_visible(False)

fig,ax=plt.subplots(figsize=(8.7,3))
ax.plot(hours,means,"-o",c=BLUE,lw=2.1,markersize=3.5)
axis(ax,"서울 공공자전거 · 정상 운영 자료의 시간대별 평균")
for t in (4,8,18):
    ax.annotate(f"{t}시 {means[t]:,.1f}건",(t,means[t]),
                xytext=(t-2,means[t]+160 if t!=18 else means[t]-200),
                fontsize=8.5)
fig.tight_layout();fig.savefig(FIG/"hourly.png",dpi=190);plt.close(fig)

fig,ax=plt.subplots(figsize=(8.7,3))
labels=["평균 예측","24시간","24+12시간","주중·주말"]
mae=np.array([testm.loc[k,"mae"] for k in mnames])
rmse=np.array([testm.loc[k,"rmse"] for k in mnames])
x=np.arange(4);w=.33
ax.bar(x-w/2,mae,w,label="MAE",color=BLUE)
ax.bar(x+w/2,rmse,w,label="RMSE",color=TEAL)
for i in range(4):
    ax.text(i-w/2,mae[i]+7,f"{mae[i]:.1f}",fontsize=8,ha="center")
    ax.text(i+w/2,rmse[i]+7,f"{rmse[i]:.1f}",fontsize=8,ha="center")
ax.set_xticks(x,labels);ax.set_ylim(0,700)
ax.set_ylabel("대여 건수/시간");ax.set_title("테스트 자료 1,512개에 대한 오차 비교")
ax.legend(loc="upper right");ax.grid(axis="y",alpha=.18)
fig.tight_layout();fig.savefig(FIG/"errors.png",dpi=190);plt.close(fig)

fig,ax=plt.subplots(figsize=(8.7,3))
actual=examples["actual"].to_numpy()
ax.plot(hours,actual,"o-",c=NAVY,lw=2.1,markersize=3.4,label="실제")
ax.plot(hours,examples["m1_24h"],"--",c=ORANGE,lw=1.8,label="모델 1")
ax.plot(hours,examples["m2_24h_12h"],"-",c=BLUE,lw=1.8,label="모델 2")
ax.plot(hours,examples["m3_with_weekend"],"-",c=TEAL,lw=1.8,label="모델 3")
axis(ax,"테스트 날짜 예시: 2017년 12월 2일")
ax.legend(ncol=4,loc="upper left",fontsize=8)
fig.tight_layout();fig.savefig(FIG/"example_day.png",dpi=190);plt.close(fig)

D=Document()
sec=D.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.45)
sec.left_margin=Cm(1.65);sec.right_margin=Cm(1.65)
sec.header_distance=Cm(.75);sec.footer_distance=Cm(.75)
styles=D.styles
norm=styles["Normal"]
norm.font.name=FONT;norm.font.size=Pt(9.2)
norm.font.color.rgb=RGBColor(33,48,62)
norm._element.rPr.rFonts.set(qn("w:eastAsia"),FONT)
norm.paragraph_format.line_spacing=1.2
norm.paragraph_format.space_after=Pt(4)
for name,size in (("Title",17),("Heading 1",13),("Heading 2",10)):
    st=styles[name];st.font.name=FONT;st.font.size=Pt(size)
    st.font.color.rgb=RGBColor(23,50,77);st.font.bold=True
    st._element.rPr.rFonts.set(qn("w:eastAsia"),FONT)
    st.paragraph_format.space_before=Pt(6);st.paragraph_format.space_after=Pt(4)
caption_style=styles.add_style("ReportCaption",WD_STYLE_TYPE.PARAGRAPH)
caption_style.font.name=FONT;caption_style.font.size=Pt(8.0)
caption_style.font.color.rgb=RGBColor(94,112,131)
caption_style._element.rPr.rFonts.set(qn("w:eastAsia"),FONT)
cap=D.add_paragraph
hdr=sec.header.paragraphs[0];hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hdr.add_run("수학적 모델링 탐구 보고서  |  서울 공공자전거")
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run("UCI Seoul Bike Sharing Demand · ")
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");foot._p.append(field)

def p(s):
    return D.add_paragraph(s)
def sub(s):
    D.add_paragraph(s,style="Heading 2")
def pic(name,width=16.6):
    par=D.add_paragraph();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after=Pt(1)
    par.add_run().add_picture(str(FIG/name),width=Cm(width))
def caption(s):
    D.add_paragraph(s,style="ReportCaption")
def table(head,body):
    t=D.add_table(rows=1,cols=len(head))
    t.style="Light Shading Accent 1";t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(head):t.cell(0,i).text=str(h)
    for row in body:
        cells=t.add_row().cells
        for i,v in enumerate(row):cells[i].text=str(v)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                paragraph.paragraph_format.space_before=Pt(2)
                for run in paragraph.runs:
                    run.font.size=Pt(8.25);run.font.name=FONT
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
    D.add_paragraph().paragraph_format.space_after=Pt(0)
def page():
    D.add_page_break()

D.add_paragraph("서울시 공공자전거 대여량의 시간적 주기성 분석","Title")
p("삼각함수 모델과 최소제곱법을 이용한 대여량 예측 · 실제 데이터 기반 탐구")
table(["항목","내용"],[
    ("학번 / 이름","2211 / 이은세"),
    ("탐구 주제","서울시 공공자전거 대여량의 시간적 주기성과 삼각함수 모델의 예측 정확도 비교"),
    ("데이터 출처","UCI Machine Learning Repository · Seoul Bike Sharing Demand (2020)\nDOI: https://doi.org/10.24432/C5F62R"),
    ("데이터 표본 수(n)","원본 n=8,760 / 미운영 295개 제외, 분석 n=8,465")])
D.add_paragraph("1. 수집한 데이터 정리 및 시각화","Heading 1")
p("2017년 12월 1일부터 2018년 11월 30일까지의 서울 공공자전거 시간별 대여 기록을 사용하였다. 원본은 14개 변수와 8,760개 시간별 관측값으로 구성된다. 이 중 서비스 미운영(Functioning Day=No) 295행을 제외한 8,465행을 분석하였다. 날짜를 일/월/년 형식으로 해석하고 결측·중복 및 자료형을 검사하였다. 시간 t(0~23시)를 독립변수, 시간당 대여량 y를 종속변수로 정하였다.")
pic("hourly.png")
caption("그림 1. 정상 운영 기록에서 계산한 시간대별 평균 대여 건수.")
table(["시각","04시","08시","18시"],[["평균 대여량(건)",f"{means[4]:.1f}",f"{means[8]:,.1f}",f"{means[18]:,.1f}"]])
p("새벽 4시에는 약 137.5건이지만 오전 8시와 오후 6시에는 각각 약 1,050.2건, 1,554.0건으로 증가한다. 하루에 두 개의 상승 구간이 있으므로 24시간 단일 주기함수보다 12시간 성분을 추가했을 때 예측 오차가 감소하는지 검증하였다.")

page()
D.add_paragraph("2. 수학적 모델링 및 예측값 도출","Heading 1")
p("같은 날짜의 기록이 서로 다른 집합에 중복되지 않도록 날짜 단위로 나누고 계절과 주말 여부의 비율을 고려하였다. 학습 243일(5,832행), 검증 47일(1,121행), 테스트 63일(1,512행)이며, 계수 추정에는 학습 자료만 사용하였다.")
sub("① 모델 1: 24시간 주기")
pic("eq1.png",15.4)
sub("② 모델 2: 24시간과 12시간 주기")
pic("eq2.png",15.4)
p("12시간 주기의 사인·코사인항을 추가해 하루 두 번 증가하는 현상을 나타낸다. 삼각함수는 시간 t에 관해서는 비선형이지만 추정 계수에 관해서는 선형이다.")
sub("③ 모델 3: 주중·주말을 구분")
p("주중 w=0, 주말 w=1로 놓고 모델 2에 w·g(t)를 더했다. g(t)는 상수항과 24시간 및 12시간 사인·코사인 항을 포함하므로 주중·주말의 서로 다른 파형을 표현한다.")
sub("④ 최소제곱법과 실제 추정 계수")
p("각 t에서 계산한 사인·코사인값을 행렬 X의 한 행으로 놓아 예측값을 Xβ로 나타냈다. 잔차제곱합을 최소화하는 계수를 구하며 정규방정식 XᵀXβ=Xᵀy를 이용해 이론적 과정을 설명할 수 있다. 실제 계산에는 수치적으로 안정적인 numpy.linalg.lstsq를 사용했다.")
pic("least.png",13.8)
coefs=c["harmonic_24h_12h"]
table(["모델 2 계수","a₀","a₁","b₁","a₂","b₂"],
    [["추정값",*map(lambda x:f"{x:.3f}",coefs)]])
p("모델 2 추정식은 ŷ₂(t)=722.781−403.395sin(2πt/24)−91.721cos(2πt/24)−183.426sin(4πt/24)−116.934cos(4πt/24)이다. 24시간 성분 진폭은 약 413.69, 12시간 성분 진폭은 약 217.53이고, 진폭비는 약 0.526이다.")
sub("⑤ 예측값 대입 계산")
p("미학습 테스트 날짜인 2017년 12월 2일 08시에 실제 대여량은 219건이었다. 모델 2에 t=8을 대입하면 636.609건으로 예측되며 절대오차는 |219−636.609|=417.609건이다. 같은 시각 모델 1은 419.291건(오차 200.291건), 모델 3은 347.717건(오차 128.717건)을 예측했다.")

page()
D.add_paragraph("3. 오차 분석, 모델 한계점 및 개선 방안","Heading 1")
p("평균적인 오차는 MAE(평균절대오차), 큰 오차에 민감한 지표는 RMSE(평균제곱근오차), 전체 변동의 설명 정도는 결정계수 R²로 비교하였다. 모델을 평가할 때는 학습에 사용하지 않은 1,512개 관측값만 사용하였다.")
pic("error.png",14.7)
table(["모델","테스트 MAE","테스트 RMSE","R²"],[
 [label,f"{testm.loc[key,'mae']:.3f}",f"{testm.loc[key,'rmse']:.3f}",f"{testm.loc[key,'r2']:.4f}"]
 for key,label in zip(mnames,["평균만 예측","M1 · 24시간","M2 · 24+12시간","M3 · 주중·주말"])])
pic("errors.png")
caption("그림 2. 테스트 데이터의 모델별 MAE와 RMSE 비교(단위: 건/시간).")
before=testm.loc["harmonic_24h","mae"]
after=testm.loc["harmonic_24h_12h","mae"]
p(f"M1에서 M2로 확장하자 MAE는 ({before:.3f}−{after:.3f})/{before:.3f}×100={100*(before-after)/before:.2f}% 감소했다. M3에서는 MAE가 405.584건/시간으로 더 줄었다. 다만 R²는 약 0.304에 불과해 일주기 요인만으로는 대여량의 대부분을 설명할 수 없다.")
sub("한계와 오차 발생 원인")
p("같은 시각이라도 강수·기온·계절·공휴일에 따라 대여량이 달라지지만 이번 모델은 이를 설명변수로 넣지 않았다. 2017년 12월 2일 18시 실제 대여량은 489건인 반면 모델 2 예측값은 1,243.109건으로, 절대오차는 754.109건이었다.")
sub("구체적인 수학적 개선 방안")
p("첫째, 강수량 R·기온 T·공휴일 H를 추가하여 ŷ=f(t,w)+γ₁R+γ₂T+γ₃H 형태의 다중회귀를 적용한다. 둘째, 계절 더미변수와 사인·코사인 항의 상호작용을 추가하여 계절별 진폭·위상 차이를 표현한다. 셋째, 과적합을 막기 위해 릿지회귀 min(||Xβ−y||²+λ||β||²)를 검토하고 λ를 검증 데이터에서 선택한다.")
p("자료의 시간 구간이 1년뿐이므로 이 결과는 그 기간 안의 다른 날짜에 대한 예측 성능이지 다음 연도의 미래 예측 성능이 아니다. 또한 최소제곱 회귀는 음수 대여량을 예측할 수 있으므로 실제 수요 예측 서비스에는 추가 제약이 필요하다.")

page()
D.add_paragraph("4. 탐구 성찰 (자기 평가)","Heading 1")
p("처음에는 시간별 대여량을 24시간 주기 하나로 표현할 수 있을 것으로 예상했다. 그러나 실제 시간대별 평균에서 오전 8시와 오후 6시에 서로 다른 대여량 증가 구간이 관측되어 12시간 성분이 필요할 수 있다는 가설을 세웠다. 두 조화 성분을 가진 모델을 학습·테스트 데이터에 적용한 결과 MAE가 약 3.65% 감소했다. 단순히 그래프가 비슷해 보이는 것을 넘어서 실제 예측값과 수치 지표를 비교하는 것이 중요함을 확인했다.")
p("동시에 모델이 복잡해졌다고 모든 날짜에서 정확해지는 것은 아니었다. 2017년 12월 2일 08시에는 모델 2가 모델 1보다 더 크게 틀렸다. 개별 예시와 전체 테스트 평균을 구분하고, 같은 날짜의 관측값이 학습과 테스트에 중복되지 않도록 분할해야 한다는 점을 검토했다.")
p("수학적으로는 삼각함수를 사용한 함수도 계수에 관해서는 선형이므로 선형대수의 최소제곱법으로 해를 구할 수 있다는 점을 이해했다. 최종 모델의 R²가 약 0.304로 낮은 사실은 주기성만으로 복잡한 현실의 대여량을 충분히 설명하지 못한다는 한계를 드러낸다. 후속 탐구에서는 기상 조건과 계절을 함께 설명변수로 사용한 뒤 같은 평가 조건에서 오차가 실제로 줄어드는지 확인하고자 한다.")
sub("자료 출처와 재현 정보")
p("UCI Machine Learning Repository. (2020). Seoul Bike Sharing Demand [Dataset]. https://doi.org/10.24432/C5F62R (CC BY 4.0).")
p("분석 코드와 원자료: https://github.com/aether0xff-afk/LinearAlgebra_Seoul_bike")
pic("example_day.png",15.9)
caption("그림 3. 테스트 날짜의 실제 시간별 대여량과 모델별 예측 곡선.")

target=OUT/"서울_공공자전거_수학적모델링_탐구보고서.docx"
D.save(target)
print(f"Generated {target} ({target.stat().st_size} bytes)")
