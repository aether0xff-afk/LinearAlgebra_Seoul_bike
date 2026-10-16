"""Reformat generated report using the original classroom form's table layout.

Keep the source document's graphs, equations and text; regroup them into the
four required numbered fields after a two-column identification table.
Run after build_submission.py and before DOCX-to-PDF conversion.
"""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

STEM="서울_공공자전거_수학적모델링_탐구보고서"
PATH=Path("submission")/(STEM+".docx")
doc=Document(PATH)
font="NanumGothic"
items={1:[],2:[],3:[],4:[]}
heads={
    1:"1. 수집한 데이터 정리 및 시각화",
    2:"2. 수학적 모델링 및 예측값 도출",
    3:"3. 오차 분석, 모델 한계점 및 개선 방안",
    4:"4. 탐구 성찰 (자기 평가)",
}
subtitles={
    1:"표, 그래프 등을 활용하여 데이터의 경향성 파악",
    2:"함수식 또는 통계 모델 수립 과정 및 이를 통한 결과 예측",
    3:"실제 데이터와 예측값의 오차 수식 계산 및 원인 분석, 모델의 한계점 비판 및 구체적인 수학적 개선 대안 제시",
    4:"탐구 과정과 결과에 대한 반성, 자신의 비판적 사고 과정 점검 및 보완점",
}

body=doc.element.body
current=None
for element in list(body):
    if element.tag==qn("w:sectPr"):
        continue
    content="".join(t.text or "" for t in element.iter(qn("w:t")))
    found=next((i for i,h in heads.items() if content.strip().startswith(h)),None) if element.tag==qn("w:p") else None
    if found is not None:
        current=found
    elif current is not None:
        is_page_break=(element.tag==qn("w:p") and
                       element.find(".//"+qn("w:br")) is not None and
                       not content.strip())
        if not is_page_break:
            items[current].append(element)
    body.remove(element)
assert all(len(items[i])>=3 for i in range(1,5)), "Missing original report section"

sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(.95);sec.bottom_margin=Cm(1.02)
sec.left_margin=Cm(1.29);sec.right_margin=Cm(1.29)
for part in (sec.header,sec.footer):
    for p in part.paragraphs:
        p.clear()

def fmt(run,size=9,bold=False,color="111111"):
    run.font.name=font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),font)
    run.font.size=Pt(size);run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    return run

def style_cell(cell,fill=None):
    tcPr=cell._tc.get_or_add_tcPr()
    if fill:
        element=OxmlElement("w:shd");element.set(qn("w:fill"),fill)
        tcPr.append(element)
    borders=OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right"):
        el=OxmlElement("w:"+side)
        el.set(qn("w:val"),"single");el.set(qn("w:sz"),"6")
        el.set(qn("w:color"),"353535");borders.append(el)
    tcPr.append(borders)
    mar=OxmlElement("w:tcMar")
    for side,width in (("top",100),("start",160),("bottom",100),("end",150)):
        e=OxmlElement("w:"+side)
        e.set(qn("w:w"),str(width));e.set(qn("w:type"),"dxa")
        mar.append(e)
    tcPr.append(mar)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def widths(table):
    table.autofit=False
    table.columns[0].width=Cm(4.65)
    table.columns[1].width=Cm(13.78)

def section(table,i):
    hdrow=table.add_row()
    header=hdrow.cells[0].merge(hdrow.cells[1])
    style_cell(header,"F2F3F5")
    p=header.paragraphs[0];p.paragraph_format.space_after=Pt(3)
    fmt(p.add_run(heads[i]),10.1,True)
    sub=header.add_paragraph();sub.paragraph_format.space_after=Pt(0)
    fmt(sub.add_run("("+subtitles[i]+")"),8.1,color="5F6977")
    contentrow=table.add_row()
    cell=contentrow.cells[0].merge(contentrow.cells[1])
    style_cell(cell)
    cell._tc.remove(cell.paragraphs[0]._p)
    for element in items[i]:
        cell._tc.append(element)
    if cell._tc[-1].tag!=qn("w:p"):
        cell.add_paragraph()

p=doc.add_paragraph()
p.paragraph_format.space_after=Pt(8)
fmt(p.add_run("[수행평가 2]  실생활 데이터 기반 수학적 모델링 및 한계점 분석"),11)

meta=doc.add_table(rows=4,cols=2)
widths(meta)
values=[
    ("학번 / 이름","2211 / 이은세"),
    ("탐구 주제","서울시 공공자전거 대여량의 시간적 주기성과 삼각함수 모델의 예측 정확도 비교"),
    ("데이터 출처\n(통계청, 공공데이터포털 등 정확히 기재)",
     "UCI Machine Learning Repository, Seoul Bike Sharing Demand (2020)\nDOI: 10.24432/C5F62R\nhttps://archive.ics.uci.edu/dataset/560/seoul+bike+sharing+demand"),
    ("데이터 표본 수 (n)\n(정량평가: 30개 이상 권장)",
     "n = 8,760개 (원본)\n실제 분석 n = 8,465개 (서비스 미운영 295개 제외)"),
]
for i,(label,value) in enumerate(values):
    for col,content in ((0,label),(1,value)):
        cell=meta.cell(i,col)
        style_cell(cell,"E4EDF9" if col==0 else None)
        p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1.35
        for j,line in enumerate(content.split("\n")):
            if j:
                p.add_run("\n")
            fmt(p.add_run(line),9.1 if j==0 else 8.25,
                bold=(col==0 and j==0),
                color=("BB1919" if (i==3 and col==0) else "202020"))
section(meta,1)
for i in (2,3,4):
    doc.add_page_break()
    table=doc.add_table(rows=0,cols=2)
    widths(table)
    section(table,i)

doc.save(PATH)
print("Applied original assessment form to",PATH)
